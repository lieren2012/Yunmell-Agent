"""Unit tests for in-process document indexing jobs."""

from __future__ import annotations

import asyncio
import threading
from pathlib import Path
from types import SimpleNamespace

import pytest

from octop.infra.db.migrate import run_migrations
from octop.infra.db.pool import SqlitePool
from octop.infra.db.repos.knowledge import KnowledgeRepo
from octop.infra.db.repos.settings import SettingsRepo
from octop.infra.db.repos.users import UserRepo
from octop.infra.knowledge import jobs
from octop.infra.knowledge.files import write_document
from octop.infra.knowledge.index import KnowledgeIndex
from octop.infra.knowledge.service import KnowledgeService


def test_enqueue_index_document_limits_concurrency(monkeypatch: pytest.MonkeyPatch) -> None:
    jobs.reset_index_semaphore_for_tests()

    lock = threading.Lock()
    active = 0
    peak = 0
    entered = threading.Event()
    hold = threading.Event()

    def blocking(_services: object, _kb: str, _doc: str) -> None:
        nonlocal active, peak
        with lock:
            active += 1
            peak = max(peak, active)
        entered.set()
        assert hold.wait(timeout=2)
        with lock:
            active -= 1

    monkeypatch.setattr(jobs, "process_document", blocking)

    async def run() -> None:
        first = jobs.enqueue_index_document(object(), "kb", "d1")
        second = jobs.enqueue_index_document(object(), "kb", "d2")
        for _ in range(100):
            if entered.is_set():
                break
            await asyncio.sleep(0.01)
        await asyncio.sleep(0.05)
        assert peak == 1
        hold.set()
        await asyncio.gather(first, second)

    asyncio.run(run())
    assert peak == 1


def test_process_document_marks_empty_extract_as_failed(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OCTOP_HOME", str(tmp_path / "home"))
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(
        owner_user_id=owner_id, name="Docs", embedding_model="BAAI/bge-small-zh-v1.5"
    )
    doc = repo.create_document(
        kb_id=kb.id, filename="blank.md", content_type="text/markdown", byte_size=0
    )
    write_document(kb.id, doc.id, doc.filename, b"   \n")
    services = SimpleNamespace(knowledge_repo=repo, settings_repo=SettingsRepo(pool))
    monkeypatch.setattr(jobs, "assert_knowledge_usable", lambda *_args: None)
    monkeypatch.setattr(
        jobs, "embed_knowledge_texts", lambda *_a, **_k: (_ for _ in ()).throw(AssertionError())
    )

    with pytest.raises(ValueError, match="no extractable text"):
        jobs.process_document(services, kb.id, doc.id)

    updated = repo.get_document(doc.id)
    assert updated is not None
    assert updated.status == "failed"
    assert updated.chunk_count == 0


@pytest.mark.parametrize("filename", ["blank.xlsx", "blank.pptx"])
def test_process_document_marks_empty_office_extract_with_actionable_reason(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, filename: str
) -> None:
    from openpyxl import Workbook
    from pptx import Presentation

    monkeypatch.setenv("OCTOP_HOME", str(tmp_path / "home"))
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(owner_user_id=owner_id, name="Docs", embedding_model="model")
    doc = repo.create_document(
        kb_id=kb.id,
        filename=filename,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        if filename.endswith(".xlsx")
        else "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        byte_size=0,
    )
    path = tmp_path / "office.bin"
    if filename.endswith(".xlsx"):
        Workbook().save(path)
    else:
        Presentation().save(path)
    write_document(kb.id, doc.id, filename, path.read_bytes())
    services = SimpleNamespace(knowledge_repo=repo, settings_repo=SettingsRepo(pool))
    monkeypatch.setattr(jobs, "optional_ocr_extractor", lambda *_a, **_k: None)
    monkeypatch.setattr(
        jobs,
        "embed_knowledge_texts",
        lambda *_a, **_k: (_ for _ in ()).throw(AssertionError()),
    )

    with pytest.raises(ValueError, match="文件中没有可提取的文本"):
        jobs.process_document(services, kb.id, doc.id)

    updated = repo.get_document(doc.id)
    assert updated is not None
    assert updated.status == "failed"
    assert updated.error_message
    assert ("Excel" in updated.error_message) or ("PowerPoint" in updated.error_message)


def test_failed_xlsx_can_be_reindexed_after_reupload(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from openpyxl import Workbook

    monkeypatch.setenv("OCTOP_HOME", str(tmp_path / "home"))
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(owner_user_id=owner_id, name="Docs", embedding_model="model")
    doc = repo.create_document(
        kb_id=kb.id,
        filename="report.xlsx",
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        byte_size=4,
    )
    write_document(kb.id, doc.id, doc.filename, b"broken")
    services = SimpleNamespace(knowledge_repo=repo, settings_repo=SettingsRepo(pool))
    monkeypatch.setattr(jobs, "optional_ocr_extractor", lambda *_a, **_k: None)
    with pytest.raises(ValueError, match="Excel 文件无法解析"):
        jobs.process_document(services, kb.id, doc.id)
    failed = repo.get_document(doc.id)
    assert failed is not None and failed.status == "failed"

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.append(["恢复后的内容"])
    valid_path = tmp_path / "valid.xlsx"
    workbook.save(valid_path)
    write_document(kb.id, doc.id, doc.filename, valid_path.read_bytes())

    refreshed = KnowledgeService(services).reindex_document(kb.id, doc.id, actor_user_id=owner_id)
    assert refreshed.status == "pending"
    assert refreshed.error_message == ""
    monkeypatch.setattr(
        jobs,
        "embed_knowledge_texts",
        lambda _services, texts, **_kwargs: [[1.0, 0.0] for _ in texts],
    )
    jobs.process_document(services, kb.id, doc.id)
    recovered = repo.get_document(doc.id)
    assert recovered is not None
    assert recovered.status == "ready"
    assert recovered.error_message == ""
    assert recovered.chunk_count == 1


def test_process_document_indexes_chunks_and_marks_ready(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OCTOP_HOME", str(tmp_path / "home"))
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(
        owner_user_id=owner_id, name="Docs", embedding_model="BAAI/bge-small-zh-v1.5"
    )
    doc = repo.create_document(
        kb_id=kb.id, filename="notes.md", content_type="text/markdown", byte_size=5
    )
    write_document(kb.id, doc.id, doc.filename, b"hello")
    services = SimpleNamespace(knowledge_repo=repo, settings_repo=SettingsRepo(pool))
    monkeypatch.setattr(jobs, "assert_knowledge_usable", lambda *_args: None)
    monkeypatch.setattr(
        jobs, "embed_knowledge_texts", lambda _services, texts: [[1.0, 0.0] for _ in texts]
    )

    jobs.process_document(services, kb.id, doc.id)

    updated = repo.get_document(doc.id)
    assert updated is not None
    assert updated.status == "ready"
    assert updated.chunk_count == 1
    assert KnowledgeIndex(kb.id).search([1.0, 0.0], 1)[0].doc_id == doc.id


def test_process_document_indexes_ocr_text(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("OCTOP_HOME", str(tmp_path / "home"))
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(owner_user_id=owner_id, name="Scans", embedding_model="model")
    doc = repo.create_document(
        kb_id=kb.id, filename="scan.png", content_type="image/png", byte_size=5
    )
    write_document(kb.id, doc.id, doc.filename, b"image")
    services = SimpleNamespace(knowledge_repo=repo, settings_repo=SettingsRepo(pool))
    monkeypatch.setattr(jobs, "assert_knowledge_usable", lambda *_args: None)
    monkeypatch.setattr(jobs, "optional_ocr_extractor", lambda _services: lambda _path: "发票")
    monkeypatch.setattr(jobs, "embed_knowledge_texts", lambda _services, _texts: [[1.0, 0.0]])

    jobs.process_document(services, kb.id, doc.id)

    updated = repo.get_document(doc.id)
    assert updated is not None
    assert updated.status == "ready"
    assert KnowledgeIndex(kb.id).search([1.0, 0.0], 1)[0].text == "发票"


def test_reindex_all_updates_base_model_marks_documents_pending_and_enqueues(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(owner_user_id=owner_id, name="Docs", embedding_model="old-model")
    doc = repo.create_document(
        kb_id=kb.id, filename="notes.md", content_type="text/markdown", byte_size=5
    )
    repo.update_document(doc.id, status="ready", chunk_count=1)
    services = SimpleNamespace(knowledge_repo=repo)
    enqueued: list[tuple[str, str]] = []
    monkeypatch.setattr(
        jobs,
        "enqueue_index_document",
        lambda _services, kb_id, doc_id: enqueued.append((kb_id, doc_id)),
    )

    jobs.reindex_all_documents(services, "new-model")

    assert repo.get_base(kb.id).embedding_model == "new-model"
    assert repo.get_base(kb.id).embedding_dim == 0
    assert repo.get_document(doc.id).status == "pending"
    assert enqueued == [(kb.id, doc.id)]


def test_reindex_shared_only_keeps_personal_model_selection(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    shared = repo.create_base(owner_user_id=owner_id, name="Shared", shared=True)
    personal = repo.create_base(
        owner_user_id=owner_id, name="Personal", embedding_model="local-model"
    )
    shared_doc = repo.create_document(
        kb_id=shared.id, filename="shared.md", content_type="text/markdown", byte_size=5
    )
    personal_doc = repo.create_document(
        kb_id=personal.id, filename="personal.md", content_type="text/markdown", byte_size=5
    )
    repo.update_document(personal_doc.id, status="ready", chunk_count=1)
    services = SimpleNamespace(knowledge_repo=repo)
    enqueued: list[tuple[str, str]] = []
    monkeypatch.setattr(
        jobs,
        "enqueue_index_document",
        lambda _services, kb_id, doc_id: enqueued.append((kb_id, doc_id)),
    )

    jobs.reindex_all_documents(services, "enterprise-model", shared_only=True)

    assert repo.get_base(shared.id).embedding_model == "enterprise-model"
    assert repo.get_base(personal.id).embedding_model == "local-model"
    assert repo.get_document(shared_doc.id).status == "pending"
    assert repo.get_document(personal_doc.id).status == "ready"
    assert enqueued == [(shared.id, shared_doc.id)]


def test_reindex_personal_only_keeps_shared_model_selection(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    shared = repo.create_base(
        owner_user_id=owner_id, name="Shared", shared=True, embedding_model="enterprise-model"
    )
    personal = repo.create_base(
        owner_user_id=owner_id, name="Personal", embedding_model="old-personal-model"
    )
    shared_doc = repo.create_document(
        kb_id=shared.id, filename="shared.md", content_type="text/markdown", byte_size=5
    )
    personal_doc = repo.create_document(
        kb_id=personal.id, filename="personal.md", content_type="text/markdown", byte_size=5
    )
    repo.update_document(shared_doc.id, status="ready", chunk_count=1)
    repo.update_document(personal_doc.id, status="ready", chunk_count=1)
    services = SimpleNamespace(knowledge_repo=repo)
    enqueued: list[tuple[str, str]] = []
    monkeypatch.setattr(
        jobs,
        "enqueue_index_document",
        lambda _services, kb_id, doc_id: enqueued.append((kb_id, doc_id)),
    )

    jobs.reindex_all_documents(services, "new-personal-model", personal_only=True)

    assert repo.get_base(shared.id).embedding_model == "enterprise-model"
    assert repo.get_base(personal.id).embedding_model == "new-personal-model"
    assert repo.get_document(shared_doc.id).status == "ready"
    assert repo.get_document(personal_doc.id).status == "pending"
    assert enqueued == [(personal.id, personal_doc.id)]


def test_resume_pending_index_jobs_resets_processing_and_enqueues(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(owner_user_id=owner_id, name="Docs")
    doc = repo.create_document(
        kb_id=kb.id, filename="notes.md", content_type="text/markdown", byte_size=5
    )
    repo.update_document(doc.id, status="processing")
    services = SimpleNamespace(knowledge_repo=repo)
    enqueued: list[tuple[str, str]] = []
    monkeypatch.setattr(
        jobs,
        "enqueue_index_document",
        lambda _services, kb_id, doc_id: enqueued.append((kb_id, doc_id)),
    )

    jobs.resume_pending_index_jobs(services)

    assert repo.get_document(doc.id).status == "pending"
    assert enqueued == [(kb.id, doc.id)]


def test_retry_failed_index_jobs_requeues_only_failed_documents(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A document that failed before a model was usable must heal on retry."""
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    owner_id = UserRepo(pool).create(username="owner", password_hash="h", role="user")
    kb = repo.create_base(owner_user_id=owner_id, name="Docs")
    failed = repo.create_document(
        kb_id=kb.id, filename="broken.xlsx", content_type="text/plain", byte_size=5
    )
    ready = repo.create_document(
        kb_id=kb.id, filename="ok.md", content_type="text/markdown", byte_size=5
    )
    repo.update_document(failed.id, status="failed", error_message="no embedding", chunk_count=0)
    repo.update_document(ready.id, status="ready", chunk_count=2)
    services = SimpleNamespace(knowledge_repo=repo)
    enqueued: list[tuple[str, str]] = []
    monkeypatch.setattr(
        jobs,
        "enqueue_index_document",
        lambda _services, kb_id, doc_id: enqueued.append((kb_id, doc_id)),
    )

    count = jobs.retry_failed_index_jobs(services)

    assert count == 1
    refreshed = repo.get_document(failed.id)
    assert refreshed.status == "pending"
    assert refreshed.error_message == ""
    assert repo.get_document(ready.id).status == "ready"
    assert enqueued == [(kb.id, failed.id)]


def test_retry_failed_respects_personal_scope(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    pool = SqlitePool(tmp_path / "octop.db")
    run_migrations(pool)
    repo = KnowledgeRepo(pool)
    users = UserRepo(pool)
    owner_id = users.create(username="owner", password_hash="h", role="user")
    other_id = users.create(username="other", password_hash="h", role="user")
    mine = repo.create_base(owner_user_id=owner_id, name="Mine")
    theirs = repo.create_base(owner_user_id=other_id, name="Theirs")
    shared = repo.create_base(owner_user_id=owner_id, name="Shared", shared=True)
    docs = {}
    for name, base in (("mine", mine), ("theirs", theirs), ("shared", shared)):
        docs[name] = repo.create_document(
            kb_id=base.id, filename=f"{name}.md", content_type="text/markdown", byte_size=5
        )
        repo.update_document(docs[name].id, status="failed", error_message="boom", chunk_count=0)
    services = SimpleNamespace(knowledge_repo=repo)
    monkeypatch.setattr(jobs, "enqueue_index_document", lambda *args: None)

    count = jobs.retry_failed_index_jobs(services, personal_only=True, personal_owner_id=owner_id)

    assert count == 1
    assert repo.get_document(docs["mine"].id).status == "pending"
    assert repo.get_document(docs["theirs"].id).status == "failed"
    assert repo.get_document(docs["shared"].id).status == "failed"
